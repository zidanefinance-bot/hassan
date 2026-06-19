"""
Real Gmail data se Quotation_Tracker.xlsx build karta hai
"""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

EXCEL_FILE = "Quotation_Tracker.xlsx"

HEADERS = [
    "Sr#", "Requester Name", "Email ID", "PR Number",
    "Item Description", "UOM", "Qty", "Rate (PKR)",
    "Query Date", "Query Time", "Status",
    "Quotation Sent Date", "Remarks", "Days Pending"
]

COLOR_HEADER  = "1F4E79"
COLOR_PENDING = "FFD700"
COLOR_OVERDUE = "FF4444"
COLOR_SENT    = "70AD47"
COLOR_FOLLOWUP= "FF8C00"

REAL_DATA = [
    # SR, Name, Email, PR, Item, UOM, Qty, Rate, Date, Time, Status, SentDate, Remarks
    (1, "Rajby Industries", "Usama.Nadeem@rajby.com.pk", "Mechanical Tools",
     "Tool Box 3-Step (Metal) + Ring Spanner Set + Makute Screwdriver Set (100pcs) + Electrical Tools",
     "Set", "Multiple", "",
     "11-Jun-2026", "02:58 PM", "OVERDUE",
     "", "Reminder x2 received (12-Jun & 19-Jun) — No quotation sent yet!", "8"),

    (2, "Rajby Industries", "Usama.Nadeem@rajby.com.pk", "PR# 928",
     "Wooden Patti 1 x 1 Inch x 10 Feet",
     "Feet", "10", "",
     "17-Jun-2026", "02:40 PM", "OVERDUE",
     "", "Reminder received 18-Jun — No quotation sent yet!", "2"),

    (3, "Rajby Industries", "Usama.Nadeem@rajby.com.pk", "PR# 23636",
     "HP EliteBook 850 G5/G6 — Refurbished Core i5 8th Gen, 8GB RAM, 512GB SSD, 15.6\" FHD, Backlit KB",
     "Nos", "1", "94,500",
     "18-Jun-2026", "03:36 PM", "Sent",
     "18-Jun-2026 04:22 PM", "Quoted @ Rs 94,500", ""),

    (4, "Rajby Industries", "Shoaib.Shahid@rajby.com.pk", "—",
     "Finishing WAT Dispenser — 3A-WAT202 (Automatic Tape Cutter Machine)",
     "Nos", "2", "",
     "19-Jun-2026", "11:52 AM", "Pending",
     "", "Video & pictures attached by client", "0"),

    (5, "Rajby Industries", "M.anas@rajby.com.pk", "PR-22816",
     "Stitching Machine FEEDO — JUKI MS 1261 [FOA] Misc. Regulator Controller (Meter) Complete Set",
     "Set", "4", "",
     "19-Jun-2026", "09:39 AM", "Pending",
     "", "Brand mention required. Letterhead + all taxes inclusive", "0"),

    (6, "Hunar Foundation", "wajeeha.nasir@hunarfoundation.org", "RFQ: Fans",
     "Fans (Various Sizes) — Delivery: SZAKATI Institute, Orangi Town",
     "Nos", "Multiple", "",
     "19-Jun-2026", "11:10 AM", "Sent",
     "19-Jun-2026 04:27 PM", "QT-001073 sent. Client asked for fan pictures — FOLLOW UP NEEDED", ""),
]

def build_excel():
    if os.path.exists(EXCEL_FILE):
        os.remove(EXCEL_FILE)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quotation Log"

    thin = Side(style='thin', color="AAAAAA")
    thick = Side(style='medium', color="000000")
    header_border = Border(left=thick, right=thick, top=thick, bottom=thick)
    cell_border   = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ── HEADER ROW ──
    header_fill = PatternFill(start_color=COLOR_HEADER, end_color=COLOR_HEADER, fill_type="solid")
    for col, h in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill   = header_fill
        c.font   = Font(color="FFFFFF", bold=True, size=11)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = header_border
    ws.row_dimensions[1].height = 38
    ws.freeze_panes = "A2"

    # ── DATA ROWS ──
    for row_idx, entry in enumerate(REAL_DATA, 2):
        sr, name, email, pr, item, uom, qty, rate, date, time_, status, sent_date, remarks, days = entry

        row_data = [sr, name, email, pr, item, uom, qty, rate,
                    date, time_, status, sent_date, remarks, days]

        if status == "OVERDUE":
            base_color = COLOR_OVERDUE
            status_color = COLOR_OVERDUE
            font_color = "FFFFFF"
        elif status == "Sent":
            base_color = "E8F5E0"
            status_color = COLOR_SENT
            font_color = "FFFFFF"
        elif status == "Pending":
            base_color = "FFFDE7"
            status_color = COLOR_PENDING
            font_color = "000000"
        else:
            base_color = "FFFFFF"
            status_color = COLOR_PENDING
            font_color = "000000"

        for col, val in enumerate(row_data, 1):
            c = ws.cell(row=row_idx, column=col, value=val)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = cell_border

            if col == 11:  # Status column
                c.fill = PatternFill(start_color=status_color, end_color=status_color, fill_type="solid")
                c.font = Font(bold=True, color=font_color, size=10)
            elif status == "OVERDUE":
                c.fill = PatternFill(start_color=COLOR_OVERDUE, end_color=COLOR_OVERDUE, fill_type="solid")
                c.font = Font(color="FFFFFF", size=10)
            else:
                c.fill = PatternFill(start_color=base_color, end_color=base_color, fill_type="solid")
                c.font = Font(size=10)

            if col == 13:  # Remarks — left align
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws.row_dimensions[row_idx].height = 40

    # ── COLUMN WIDTHS ──
    widths = [5, 22, 38, 16, 52, 8, 8, 12, 14, 12, 12, 22, 45, 12]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    # ── SUMMARY SECTION ──
    summary_row = len(REAL_DATA) + 3
    ws.cell(row=summary_row, column=1, value="SUMMARY").font = Font(bold=True, size=12, color="1F4E79")

    pending  = [e for e in REAL_DATA if e[10] == "Pending"]
    overdue  = [e for e in REAL_DATA if e[10] == "OVERDUE"]
    sent     = [e for e in REAL_DATA if e[10] == "Sent"]

    ws.cell(row=summary_row+1, column=1, value=f"Total Queries: {len(REAL_DATA)}")
    ws.cell(row=summary_row+2, column=1, value=f"✅ Sent: {len(sent)}")
    ws.cell(row=summary_row+2, column=1).font = Font(color="2E7D32", bold=True)
    ws.cell(row=summary_row+3, column=1, value=f"⏳ Pending: {len(pending)}")
    ws.cell(row=summary_row+3, column=1).font = Font(color="F57F17", bold=True)
    ws.cell(row=summary_row+4, column=1, value=f"🚨 Overdue: {len(overdue)}")
    ws.cell(row=summary_row+4, column=1).font = Font(color="C62828", bold=True)

    wb.save(EXCEL_FILE)
    print(f"\n✅ Excel ready: {EXCEL_FILE}")
    print(f"\n📊 Summary:")
    print(f"   Total  : {len(REAL_DATA)}")
    print(f"   Sent   : {len(sent)}")
    print(f"   Pending: {len(pending)}")
    print(f"   Overdue: {len(overdue)}")
    for e in overdue:
        print(f"   🚨 {e[1]} | {e[4][:50]} | {e[13]} days")

if __name__ == "__main__":
    build_excel()
