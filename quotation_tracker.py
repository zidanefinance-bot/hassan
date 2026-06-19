#!/usr/bin/env python3
"""
Quotation Query Tracker
- Har query ko Excel mein log karta hai
- Pending/overdue reminders highlight karta hai
- Multiple queries from same person track karta hai
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timedelta
import os

EXCEL_FILE = "Quotation_Tracker.xlsx"

HEADERS = [
    "Sr#", "Requester Name", "Email ID", "PR Number",
    "Item Description", "UOM", "Qty", "Rate (PKR)",
    "Query Date", "Query Time", "Status",
    "Quotation Sent Date", "Remarks", "Days Pending"
]

# Colors
COLOR_HEADER = "1F4E79"       # Dark blue header
COLOR_PENDING = "FFD700"       # Yellow - pending
COLOR_OVERDUE = "FF4444"       # Red - overdue (3+ days)
COLOR_SENT = "70AD47"          # Green - quotation sent
COLOR_ROW_ALT = "DEEAF1"       # Light blue alternate rows
COLOR_WHITE = "FFFFFF"

def get_or_create_workbook():
    if os.path.exists(EXCEL_FILE):
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Quotation Log"
        setup_headers(ws)
        setup_column_widths(ws)
    return wb, ws

def setup_headers(ws):
    header_fill = PatternFill(start_color=COLOR_HEADER, end_color=COLOR_HEADER, fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style='thin', color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    ws.row_dimensions[1].height = 35
    ws.freeze_panes = "A2"

def setup_column_widths(ws):
    widths = [5, 20, 28, 14, 30, 8, 8, 14, 14, 12, 14, 18, 22, 12]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

def get_next_row(ws):
    return ws.max_row + 1 if ws.max_row >= 1 else 2

def add_query(requester_name, email_id, pr_number, item_desc, uom, qty, rate=None, remarks=""):
    wb, ws = get_or_create_workbook()

    now = datetime.now()
    next_row = get_next_row(ws)

    # Auto Sr#
    sr = next_row - 1

    data = [
        sr,
        requester_name,
        email_id,
        pr_number,
        item_desc,
        uom,
        qty,
        rate if rate else "",
        now.strftime("%d-%b-%Y"),
        now.strftime("%I:%M %p"),
        "Pending",
        "",
        remarks,
        ""
    ]

    thin = Side(style='thin', color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    row_color = COLOR_ROW_ALT if sr % 2 == 0 else COLOR_WHITE

    for col, value in enumerate(data, 1):
        cell = ws.cell(row=next_row, column=col, value=value)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        # Default row fill
        if col not in [11]:  # Status column colored separately
            cell.fill = PatternFill(start_color=row_color, end_color=row_color, fill_type="solid")

    # Status cell - Pending = Yellow
    status_cell = ws.cell(row=next_row, column=11)
    status_cell.fill = PatternFill(start_color=COLOR_PENDING, end_color=COLOR_PENDING, fill_type="solid")
    status_cell.font = Font(bold=True)

    ws.row_dimensions[next_row].height = 22

    wb.save(EXCEL_FILE)
    print(f"✅ Query added | Row {next_row} | {requester_name} | {item_desc} | {now.strftime('%d-%b-%Y %I:%M %p')}")
    return next_row

def mark_quotation_sent(row_number, remarks=""):
    wb, ws = get_or_create_workbook()

    now = datetime.now()

    # Update Status
    ws.cell(row=row_number, column=11, value="Sent")
    ws.cell(row=row_number, column=11).fill = PatternFill(
        start_color=COLOR_SENT, end_color=COLOR_SENT, fill_type="solid")
    ws.cell(row=row_number, column=11).font = Font(bold=True, color="FFFFFF")

    # Update sent date
    ws.cell(row=row_number, column=12, value=now.strftime("%d-%b-%Y %I:%M %p"))

    if remarks:
        ws.cell(row=row_number, column=13, value=remarks)

    # Clear days pending
    ws.cell(row=row_number, column=14, value="Done")

    wb.save(EXCEL_FILE)
    print(f"✅ Quotation marked as SENT for Row {row_number}")

def refresh_reminders():
    """Recalculate days pending and highlight overdue rows"""
    wb, ws = get_or_create_workbook()

    now = datetime.now()
    overdue_list = []

    for row in range(2, ws.max_row + 1):
        status = ws.cell(row=row, column=11).value
        query_date_str = ws.cell(row=row, column=9).value
        requester = ws.cell(row=row, column=2).value

        if not query_date_str or not requester:
            continue

        if status == "Pending":
            try:
                query_date = datetime.strptime(str(query_date_str), "%d-%b-%Y")
                days_pending = (now - query_date).days
                ws.cell(row=row, column=14, value=days_pending)

                if days_pending >= 3:
                    # Overdue - Red highlight
                    red_fill = PatternFill(start_color=COLOR_OVERDUE, end_color=COLOR_OVERDUE, fill_type="solid")
                    for col in range(1, 15):
                        ws.cell(row=row, column=col).fill = red_fill
                    ws.cell(row=row, column=11).font = Font(bold=True, color="FFFFFF")
                    overdue_list.append({
                        "row": row,
                        "requester": requester,
                        "item": ws.cell(row=row, column=5).value,
                        "days": days_pending
                    })
                elif days_pending >= 1:
                    # Warning - Orange
                    orange_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
                    ws.cell(row=row, column=14).fill = orange_fill
            except Exception:
                pass

    wb.save(EXCEL_FILE)

    if overdue_list:
        print("\n🚨 OVERDUE REMINDERS:")
        print("-" * 60)
        for item in overdue_list:
            print(f"  Row {item['row']} | {item['requester']} | {item['item']} | {item['days']} days pending!")
        print("-" * 60)
    else:
        print("✅ No overdue quotations. All good!")

    return overdue_list

def show_summary():
    """Print summary of all pending quotations"""
    wb, ws = get_or_create_workbook()

    print("\n" + "="*70)
    print("         QUOTATION TRACKER SUMMARY")
    print("="*70)

    pending = []
    sent = []

    for row in range(2, ws.max_row + 1):
        requester = ws.cell(row=row, column=2).value
        if not requester:
            continue

        status = ws.cell(row=row, column=11).value
        entry = {
            "sr": ws.cell(row=row, column=1).value,
            "requester": requester,
            "email": ws.cell(row=row, column=3).value,
            "pr": ws.cell(row=row, column=4).value,
            "item": ws.cell(row=row, column=5).value,
            "date": ws.cell(row=row, column=9).value,
            "time": ws.cell(row=row, column=10).value,
            "days": ws.cell(row=row, column=14).value,
        }

        if status == "Pending":
            pending.append(entry)
        else:
            sent.append(entry)

    print(f"\n📋 PENDING QUOTATIONS: {len(pending)}")
    if pending:
        for e in pending:
            days_info = f" | ⏰ {e['days']} days" if e['days'] and e['days'] != "" else ""
            print(f"  [{e['sr']}] {e['requester']} | {e['item']} | {e['date']} {e['time']}{days_info}")

    print(f"\n✅ SENT QUOTATIONS: {len(sent)}")
    if sent:
        for e in sent:
            print(f"  [{e['sr']}] {e['requester']} | {e['item']}")

    print("="*70)

def interactive_menu():
    print("\n" + "="*50)
    print("   QUOTATION TRACKER")
    print("="*50)
    print("1. Add New Query")
    print("2. Mark Quotation as Sent")
    print("3. Refresh Reminders")
    print("4. Show Summary")
    print("5. Exit")
    print("-"*50)
    return input("Choose option: ").strip()

def get_input(prompt, required=True):
    while True:
        val = input(f"{prompt}: ").strip()
        if val or not required:
            return val
        print("  (Required field, please enter a value)")

def main():
    print("\n🗂️  QUOTATION TRACKER - Hassan's Business Tool")

    while True:
        choice = interactive_menu()

        if choice == "1":
            print("\n--- ADD NEW QUERY ---")
            name = get_input("Requester Name (e.g. Rajby Industries)")
            email = get_input("Email ID")
            pr = get_input("PR Number")
            item = get_input("Item Description")
            uom = get_input("UOM (e.g. Nos, Kg, Ltr)")
            qty = get_input("Quantity")
            rate = get_input("Rate (PKR) - press Enter if unknown", required=False)
            remarks = get_input("Remarks (optional)", required=False)

            row = add_query(name, email, pr, item, uom, qty, rate, remarks)
            print(f"\n✅ Query saved! Row number: {row}")
            print(f"📂 File: {EXCEL_FILE}")

        elif choice == "2":
            show_summary()
            row_num = get_input("\nEnter Row Number to mark as SENT")
            try:
                remarks = get_input("Remarks (optional)", required=False)
                mark_quotation_sent(int(row_num), remarks)
            except ValueError:
                print("❌ Invalid row number")

        elif choice == "3":
            print("\n🔄 Refreshing reminders...")
            refresh_reminders()

        elif choice == "4":
            show_summary()
            refresh_reminders()

        elif choice == "5":
            print("\nKhuda Hafiz! 👋")
            break

        else:
            print("❌ Invalid option, dobara try karein")

if __name__ == "__main__":
    main()
