#!/usr/bin/env python3
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Kashif Buying Sheet"

# Colors
HDR  = "1F4E79"
MED  = "2E75B6"
CAT1 = "D6E4F0"  # Medical
CAT2 = "D5E8D4"  # Kitchen
CAT3 = "FFF2CC"  # Plumbing
CAT4 = "F8CECC"  # Kitchen Tools
CAT5 = "E1D5E7"  # Other

thin = Side(style='thin', color="AAAAAA")
thick = Side(style='medium', color="000000")
bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
hbdr = Border(left=thick, right=thick, top=thick, bottom=thick)

def hcell(ws, row, col, val, bg=HDR, fc="FFFFFF", bold=True, size=11, center=True):
    c = ws.cell(row=row, column=col, value=val)
    c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    c.font = Font(color=fc, bold=bold, size=size)
    c.alignment = Alignment(horizontal="center" if center else "left", vertical="center", wrap_text=True)
    c.border = hbdr
    return c

def dcell(ws, row, col, val, bg="FFFFFF", bold=False, center=True):
    c = ws.cell(row=row, column=col, value=val)
    c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    c.font = Font(bold=bold, size=10)
    c.alignment = Alignment(horizontal="center" if center else "left", vertical="center", wrap_text=True)
    c.border = bdr
    return c

# ── TITLE ──
ws.merge_cells("A1:I1")
c = ws.cell(row=1, column=1, value="KASHIF BUYING SHEET — RAJBY INDUSTRIES")
c.fill = PatternFill(start_color=HDR, end_color=HDR, fill_type="solid")
c.font = Font(color="FFFFFF", bold=True, size=14)
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 35

ws.merge_cells("A2:I2")
c = ws.cell(row=2, column=1, value="Dates: 29-Jun to 06-Jul-2026 | 11 POs | Deliver To: Main Store, Plot #8 Sector 27, Karachi")
c.fill = PatternFill(start_color=MED, end_color=MED, fill_type="solid")
c.font = Font(color="FFFFFF", size=10)
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 20

# ── HEADERS ──
headers = ["Sr#", "PO Number", "Item Description", "UOM", "Qty", "Rate (PKR)", "Amount (PKR)", "Tax", "Total incl. Tax"]
for col, h in enumerate(headers, 1):
    hcell(ws, 3, col, h)
ws.row_dimensions[3].height = 30
ws.freeze_panes = "A4"

# Column widths
widths = [5, 12, 45, 8, 8, 14, 14, 10, 16]
for col, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(col)].width = w

row = 4

def cat_header(ws, row, title, color):
    ws.merge_cells(f"A{row}:I{row}")
    c = ws.cell(row=row, column=1, value=f"  {title}")
    c.fill = PatternFill(start_color=color[1:] if color.startswith('#') else color,
                         end_color=color[1:] if color.startswith('#') else color, fill_type="solid")
    c.font = Font(bold=True, size=11, color="1F4E79")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 22
    return row + 1

def add_item(ws, row, sr, po, item, uom, qty, rate, tax_pct, bg):
    amt = qty * rate
    tax = round(amt * tax_pct / 100, 2)
    total = amt + tax
    vals = [sr, po, item, uom, qty, rate, amt, tax if tax > 0 else "—", total]
    for col, v in enumerate(vals, 1):
        c = dcell(ws, row, col, v, bg=bg, center=(col != 3))
        if col == 3:
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        if col in [6,7,8,9] and isinstance(v, (int,float)):
            c.number_format = '#,##0.00'
    ws.row_dimensions[row].height = 22
    return row + 1, total

sr = 1
grand_total = 0

# ═══════════════════════════════
# CATEGORY 1: MEDICAL / FIRST AID
# ═══════════════════════════════
row = cat_header(ws, row, "🏥  MEDICAL / FIRST AID ITEMS", "D6E4F0")

medical = [
    # PO 1183 - Washing PVT
    ("1183", "Pyodine 60ml",               "PCS", 30,  225,  0),
    ("1183", "Cotton Medicated 25grm",      "PKT", 20,  20,   0),
    ("1183", "Bandages 3\"",               "PCS", 60,  25,   0),
    ("1183", "Bandages 4\"",               "PCS", 60,  30,   0),
    ("1183", "Bandages 5\"",               "PCS", 60,  34,   0),
    ("1183", "Tourniquet",                 "PCS", 20,  200,  0),
    ("1183", "Ointment (Polyfax)",         "PCS", 20,  240,  0),
    ("1183", "Betnesol 7.5ml",             "PCS", 20,  230,  0),
    ("1183", "Dettol Bottle 50ml",         "PCS", 20,  260,  0),
    ("1183", "Paratule",                   "PKT", 15,  320,  0),
    # PO 17427 - Cut to Pack PVT
    ("17427","Pyodine 60ml",               "PCS", 40,  233,  5),
    ("17427","Bandages 2\"",               "PCS", 36,  25,   5),
    ("17427","Bandages 3\"",               "PCS", 36,  34,   5),
    ("17427","Saniplast",                  "PKT", 12,  357,  5),
    ("17427","Crepe Bandage",              "PCS", 24,  49,   5),
    ("17427","Surgical Cotton",            "PCS", 24,  118,  5),
    ("17427","Poly Flex",                  "PCS", 12,  250,  5),
    ("17427","Burnol 20gm (Tube)",         "PCS", 12,  123,  5),
    ("17427","Iodex Mini Jar",             "PCS", 20,  123,  5),
    ("17427","Surgical Tape",              "PCS", 12,  93,   5),
    ("17427","Benzoin Compound Tincture",  "PCS", 24,  82,   5),
    ("17427","Dettol Bottle 50ml",         "PCS", 24,  270,  5),
    ("17427","Surgical Safety Pin",        "PKT", 4,   49,   5),
]

for item_data in medical:
    row, total = add_item(ws, row, sr, *item_data, "D6E4F0")
    grand_total += total
    sr += 1

# ═══════════════════════════════
# CATEGORY 2: KITCHEN / CROCKERY
# ═══════════════════════════════
row = cat_header(ws, row, "🍽️  KITCHEN / CROCKERY ITEMS", "D5E8D4")

kitchen = [
    ("17412","Rice Spoon",                 "PCS", 24,  50,   0),
    ("17412","Ceramic Plates White 8\"",   "PCS", 6,   295,  5),
    ("17412","Ceramic Plates White 12\"",  "PCS", 6,   300,  5),
    ("17412","Glass Jug 1.5 Ltr",         "PCS", 2,   700,  5),
    ("17412","Deli Glass 6 pcs (Set)",     "SET", 2,   800,  5),
]

for item_data in kitchen:
    row, total = add_item(ws, row, sr, *item_data, "D5E8D4")
    grand_total += total
    sr += 1

# ═══════════════════════════════
# CATEGORY 3: PLUMBING ITEMS
# ═══════════════════════════════
row = cat_header(ws, row, "🔧  PLUMBING ITEMS", "FFF2CC")

plumbing = [
    ("17186","U-PVC Solution 1/2 KG Tin",              "KG",  1,  1500, 5),
    ("17186","Muslim Shower",                           "SET", 6,  800,  5),
    ("17186","U-PVC Bib Cock 1/2\" (SCHD-80)",        "PCS", 24, 300,  5),
    ("17186","Wash Basin Flexible Drain Pipe 32mm\n(with SS Mesh/Strainer)", "PCS", 6, 800, 5),
    ("17426","Plastic Jar (Small Size)",               "PCS", 36, 70,   5),
]

for item_data in plumbing:
    row, total = add_item(ws, row, sr, *item_data, "FFF2CC")
    grand_total += total
    sr += 1

# ═══════════════════════════════
# CATEGORY 4: STEEL / TOOLS
# ═══════════════════════════════
row = cat_header(ws, row, "⚙️  STEEL / TOOLS", "F8CECC")

tools = [
    ("1163", "Kafgeer 1 KG (Steel)",       "PCS", 2,  1000, 0),
    ("1163", "Kafgeer 2 KG (Steel)",       "PCS", 3,  3000, 0),
]

for item_data in tools:
    row, total = add_item(ws, row, sr, *item_data, "F8CECC")
    grand_total += total
    sr += 1

# ═══════════════════════════════
# CATEGORY 5: JANITORIAL / CLEANING
# ═══════════════════════════════
row = cat_header(ws, row, "🧹  JANITORIAL / CLEANING ITEMS", "E2EFDA")

janitorial = [
    ("17257","Air Freshner Room Spray 300ml Topic Brand","PCS",96,  260,  5),
    ("17257","Bleach 30KG Local",                        "CAN",14,  970,  5),
    ("17257","Broom Hard",                               "PCS",200, 235,  5),
    ("17257","Broom Soft PP",                            "PCS",200, 200,  0),
    ("17257","China Brush",                              "PCS",60,  300,  5),
    ("17257","Dustbin Bag 18x24 Black",                  "KG", 50,  300,  5),
    ("17257","Dustbin Bag 30x50 Black",                  "KG", 50,  300,  5),
    ("17257","Duster Fabric",                            "PCS",96,  37,   0),
    ("17257","Glass Cleaner Glint 500ml",                "BTL",96,  250,  0),
    ("17257","Hand Brush",                               "PCS",12,  300,  0),
    ("17257","Hand Wash Lose Local",                     "KG", 200, 270,  5),
    ("17257","Handwash Lifeboy 140ml",                   "BTL",48,  300,  0),
    ("17257","Harpic 500ml",                             "BTL",96,  380,  0),
    ("17257","Lux Soap 70 Grams",                        "PCS",100, 150,  5),
    ("17257","Mop Big 18x18 Wooden",                     "PCS",100, 394,  0),
    ("17257","Mop Refill (600 Gram)",                    "PCS",200, 347,  5),
    ("17257","Mop Stick",                                "PCS",60,  473,  5),
    ("17257","Mortein Spray",                            "PCS",48,  590,  5),
    ("17257","Phenyl Zip Local 3 Ltr",                   "BTL",200, 420,  0),
    ("17257","Roomi Box",                                "CTN",2,   18000,5),
    ("17257","Safety Match 10 PCS",                      "PKT",15,  100,  5),
    ("17257","Scotch Bright",                            "PCS",24,  70,   5),
    ("17257","Scraper Plastic Supri",                    "PCS",60,  105,  5),
    ("17257","Tile Wash Sweep Local 30KG",               "CAN",7,   1180, 0),
    ("17257","Tissue Box Rose Petal Pop Up 150 Sheet",   "PKT",150, 270,  5),
    ("17257","Tissue Paper Roll",                        "PCS",300, 120,  5),
    ("17257","VIM Toilet Bleach Powder Lose Local",      "PCS",25,  180,  0),
    ("17257","Viper Large 4ft",                          "PCS",50,  579,  0),
]

for item_data in janitorial:
    row, total = add_item(ws, row, sr, *item_data, "E2EFDA")
    grand_total += total
    sr += 1

# ═══════════════════════════════
# CATEGORY 6: GROCERY / ENTERTAINMENT
# ═══════════════════════════════
row = cat_header(ws, row, "🛒  GROCERY / ENTERTAINMENT ITEMS", "FFF9C4")

grocery = [
    ("17308","Coffee (100g)",                   "PCS",3,   1600, 0),
    ("17308","Every Day Tea Whitener 2KG Pouch", "PKT",260, 4185, 0),
    ("17308","Lipton Lemon Green Tea 100 Bags",  "PKT",30,  1150, 0),
    ("17308","Lipton Tea Bag 600 Bags",          "BOX",11,  3300, 0),
    ("17308","Nescafe (200gm)",                  "PCS",18,  2800, 0),
    ("17308","Nestle Water 500ml",               "CRT",23,  680,  0),
    ("17308","Sugar 2 KGS Pack",                 "PKT",350, 305,  0),
    ("17308","Tapal Family Mixture (900gm)",      "PKT",176, 1850, 0),
]

for item_data in grocery:
    row, total = add_item(ws, row, sr, *item_data, "FFF9C4")
    grand_total += total
    sr += 1

# ═══════════════════════════════
# CATEGORY 7: PRODUCTION CONSUMABLES
# ═══════════════════════════════
row = cat_header(ws, row, "🏭  PRODUCTION CONSUMABLES", "FCE4D6")

production = [
    ("1194","Molty Foam Sheet (3'x6'x2\") 980","SHT",6,3579,0),
]

for item_data in production:
    row, total = add_item(ws, row, sr, *item_data, "FCE4D6")
    grand_total += total
    sr += 1

# ═══════════════════════════════
# CATEGORY 8: MECHANICAL ITEMS
# ═══════════════════════════════
row = cat_header(ws, row, "⚙️  MECHANICAL ITEMS", "DAEEF3")

mechanical = [
    ("406", "Ceramic Fibre Rope 10mm Dia x 30m\n(Temp 150-200°C)",    "ROL", 4, 15789, 0),
    ("76",  "Bench Vise / Wood Working Vice Set",                       "PCS", 1, 5000,  0),
    ("76",  "C-Clamp Set",                                             "PCS", 6, 1040,  0),
    ("76",  "Center Punch Set",                                        "PCS", 2, 300,   0),
    ("76",  "File Set (Rasp & File Set / Ruff File - 10mm)",           "PCS", 1, 530,   0),
    ("76",  "Flat Chisel Set",                                         "PCS", 1, 1875,  0),
    ("76",  "Hammer 1KG",                                              "PCS", 6, 480,   0),
    ("76",  "Hammer 5KG",                                              "PCS", 1, 1125,  0),
    ("76",  "Hand Level 2'",                                           "PCS", 2, 313,   0),
    ("76",  "Hole Saw Set 20mm to 50mm",                               "PCS", 1, 938,   0),
    ("76",  "Iron Chisel Rod 1'",                                      "PCS", 6, 570,   0),
    ("76",  "Iron Chisel Rod 1.5'",                                    "PCS", 6, 570,   0),
    ("76",  "Measuring Tape / Inch Tape 16FT",                         "PCS", 6, 230,   0),
    ("76",  "Nail Punch Set",                                          "PCS", 2, 875,   0),
    ("76",  "Screw Extractor Set 3mm-6mm Total Brand",                 "PCS", 1, 813,   0),
    ("76",  "Try Square 2 Ft x 6 Inch (Steel)",                       "PCS", 1, 500,   0),
    ("76",  "Try Square Steel Type",                                   "PCS", 2, 350,   0),
    ("76",  "Yato Screw Wrench Set",                                   "PCS", 1, 3000,  0),
]

for item_data in mechanical:
    row, total = add_item(ws, row, sr, *item_data, "DAEEF3")
    grand_total += total
    sr += 1

# ── GRAND TOTAL ──
row += 1
ws.merge_cells(f"A{row}:H{row}")
c = ws.cell(row=row, column=1, value="GRAND TOTAL (All POs)")
c.fill = PatternFill(start_color=HDR, end_color=HDR, fill_type="solid")
c.font = Font(color="FFFFFF", bold=True, size=12)
c.alignment = Alignment(horizontal="right", vertical="center")
c = ws.cell(row=row, column=9, value=grand_total)
c.fill = PatternFill(start_color=HDR, end_color=HDR, fill_type="solid")
c.font = Font(color="FFFFFF", bold=True, size=12)
c.alignment = Alignment(horizontal="center", vertical="center")
c.number_format = '#,##0.00'
ws.row_dimensions[row].height = 28

# ── PO SUMMARY ──
row += 2
ws.merge_cells(f"A{row}:I{row}")
c = ws.cell(row=row, column=1, value="PO SUMMARY")
c.fill = PatternFill(start_color=MED, end_color=MED, fill_type="solid")
c.font = Font(color="FFFFFF", bold=True, size=11)
c.alignment = Alignment(horizontal="center", vertical="center")
row += 1

po_summary = [
    ("PO# 1183",  "Washing PVT — First Aid Items",             "02-Jul-2026", "PKR 35,890"),
    ("PO# 17412", "Cut to Pack PVT — Kitchen/Crockery",        "02-Jul-2026", "PKR 8,099"),
    ("PO# 17426", "Cut to Pack PVT — Plastic Jar",             "02-Jul-2026", "PKR 2,646"),
    ("PO# 17427", "Cut to Pack PVT — First Aid Items",         "02-Jul-2026", "PKR 38,254"),
    ("PO# 17186", "Cut to Pack PVT — Plumbing Items",          "29-Jun-2026", "PKR 19,215"),
    ("PO# 1163",  "Washing PVT — Steel Kafgeer",               "30-Jun-2026", "PKR 11,000"),
    ("PO# 17257", "Cut to Pack PVT — Janitorial/Cleaning",     "29-Jun-2026", "PKR 760,293"),
    ("PO# 17308", "Cut to Pack PVT — Grocery/Entertainment",   "30-Jun-2026", "PKR 1,662,090"),
    ("PO# 1194",  "Washing PVT — Production Consumables",      "06-Jul-2026", "PKR 21,474"),
    ("PO# 406",   "Utilities PVT — Mechanical Items",          "03-Jul-2026", "PKR 63,156"),
    ("PO# 76",    "Denim Unit — Mechanical / Carpenter Items", "30-Jun-2026", "PKR 34,797"),
]

hcell(ws, row, 1, "PO#", MED, size=10)
ws.merge_cells(f"B{row}:F{row}")
hcell(ws, row, 2, "Description", MED, size=10)
hcell(ws, row, 7, "Date", MED, size=10)
ws.merge_cells(f"H{row}:I{row}")
hcell(ws, row, 8, "Grand Total", MED, size=10)
row += 1

for po, desc, date, total in po_summary:
    ws.cell(row=row, column=1, value=po).border = bdr
    ws.merge_cells(f"B{row}:F{row}")
    c = ws.cell(row=row, column=2, value=desc)
    c.border = bdr; c.alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=row, column=7, value=date).border = bdr
    ws.merge_cells(f"H{row}:I{row}")
    c = ws.cell(row=row, column=8, value=total)
    c.font = Font(bold=True); c.border = bdr
    c.alignment = Alignment(horizontal="center", vertical="center")
    row += 1

wb.save("Kashif_Buying_Sheet.xlsx")
print("Done!")
