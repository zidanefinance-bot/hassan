"""Rajby PR# 217: working sheet (xlsx), selling quotation PDF, internal buying vs selling PDF."""
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from items import ITEMS

MARKUP = 0.15
REF, DATE = "Rajby Industries - PR# 217", "12-Oct-2026"
RED = colors.HexColor("#B71C1C")

def r10(x):
    return int(Decimal(str(round(x, 6))).quantize(Decimal("1E1"), rounding=ROUND_HALF_UP))

for it in ITEMS:
    it["sell"] = r10(it["high"] * (1 + MARKUP))
    it["shown"] = it["brand"] + (" / Equivalent" if it["src"] == "Intl" else "")

# ---------- Excel working sheet ----------
wb = Workbook(); ws = wb.active; ws.title = "Working"
F = "Arial"; thin = Side(style="thin", color="BFBFBF"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
ws["A1"] = f"Zidane Corporation - Rate working: {REF}"; ws["A1"].font = Font(name=F, bold=True, size=13)
ws["A2"] = "Markup on Buy High:"; ws["A2"].font = Font(name=F, bold=True)
ws["D2"] = MARKUP; ws["D2"].number_format = "0%"; ws["D2"].font = Font(name=F, bold=True, color="0000FF")
ws["D2"].fill = PatternFill("solid", fgColor="FFFF00")
ws["E2"] = "<- yellow cell change karo, selling rates update ho jayenge"; ws["E2"].font = Font(name=F, italic=True, size=9)
heads = ["SR#", "Item Code", "Item Description", "UOM", "Qty", "Brand (quote)", "Buy Low", "Buy High", "Sell Rate",
         "Buy Total", "Sell Total", "Profit", "Margin", "Status V/S/N", "Src", "Source / Note"]
HR = 4
for c, h in enumerate(heads, 1):
    x = ws.cell(HR, c, h); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = PatternFill("solid", fgColor="B71C1C")
    x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); x.border = bd
r = HR + 1; first = r
for it in ITEMS:
    vals = [it["sr"], it["code"], it["desc"], it["uom"], it["qty"], it["shown"], it["low"], it["high"]]
    for c, v in enumerate(vals, 1): ws.cell(r, c, v)
    ws.cell(r, 9, f"=ROUND(H{r}*(1+$D$2),-1)")
    ws.cell(r, 10, f"=E{r}*H{r}"); ws.cell(r, 11, f"=E{r}*I{r}"); ws.cell(r, 12, f"=K{r}-J{r}")
    ws.cell(r, 13, f"=IF(K{r}=0,0,L{r}/K{r})")
    ws.cell(r, 14, it["status"]); ws.cell(r, 15, it["src"]); ws.cell(r, 16, it["note"])
    for c in range(1, 17):
        x = ws.cell(r, c); x.border = bd; x.font = Font(name=F, size=9, color="0000FF" if c in (7, 8) else "000000")
        if c in (7, 8, 9, 10, 11, 12): x.number_format = "#,##0"
        if c == 13: x.number_format = "0.0%"
    if it["status"] == "N":
        for c in range(1, 17): ws.cell(r, c).fill = PatternFill("solid", fgColor="FDECEA")
    r += 1
last = r - 1
ws.cell(r, 3, "TOTAL").font = Font(name=F, bold=True)
for c, col in ((10, "J"), (11, "K"), (12, "L")):
    x = ws.cell(r, c, f"=SUM({col}{first}:{col}{last})"); x.number_format = "#,##0"; x.font = Font(name=F, bold=True)
x = ws.cell(r, 13, f"=IF(K{r}=0,0,L{r}/K{r})"); x.number_format = "0.0%"; x.font = Font(name=F, bold=True)
ws.cell(r + 2, 1, "Status: V = exact listing, S = size/pack/brand differs (estimate in note), N = no online listing. "
        "Src: PK = Pakistani seller; Intl = UAE/UK/US/EU/IN/MY rate converted (AED x90, USD x330.5, GBP x418.5, EUR x384, PLN x90.5, MYR x78, INR x3.35).").font = Font(name=F, size=9, italic=True)
for col, w in zip("ABCDEFGHIJKLMNOP", (5, 27, 46, 9, 5, 24, 10, 10, 10, 11, 11, 10, 8, 8, 6, 80)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "D5"

q = wb.create_sheet("Quotation")
q["A1"] = "Zidane Corporation - QUOTATION"; q["A1"].font = Font(name=F, bold=True, size=14)
q["A2"] = f"Reference: {REF}    Currency: PKR    Date: {DATE}"; q["A2"].font = Font(name=F)
qh = ["SR#", "Item Code", "Item Description", "Brand", "UOM", "Qty", "Unit Rate (PKR)", "Amount (PKR)"]
for c, h in enumerate(qh, 1):
    x = q.cell(4, c, h); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = PatternFill("solid", fgColor="B71C1C"); x.border = bd
for i, it in enumerate(ITEMS):
    rr = 5 + i; wr = first + i
    for c, v in enumerate([it["sr"], it["code"], it["desc"], it["shown"], it["uom"], it["qty"]], 1): q.cell(rr, c, v)
    q.cell(rr, 7, f"=Working!I{wr}"); q.cell(rr, 8, f"=F{rr}*G{rr}")
    for c in range(1, 9):
        x = q.cell(rr, c); x.border = bd; x.font = Font(name=F, size=9, color="008000" if c == 7 else "000000")
        if c in (7, 8): x.number_format = "#,##0"
tr = 5 + len(ITEMS)
q.cell(tr, 7, "TOTAL").font = Font(name=F, bold=True)
x = q.cell(tr, 8, f"=SUM(H5:H{tr-1})"); x.number_format = "#,##0"; x.font = Font(name=F, bold=True)
for col, w in zip("ABCDEFGH", (5, 27, 50, 26, 9, 5, 15, 15)): q.column_dimensions[col].width = w
wb.save("Rajby_PR217_Working.xlsx")

# ---------- PDFs ----------
cell = ParagraphStyle("c", fontName="Helvetica", fontSize=7.4, leading=8.6)
TS = [("BACKGROUND", (0, 0), (-1, 0), RED), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
      ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8), ("FONT", (0, 1), (-1, -1), "Helvetica", 7.4),
      ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
      ("ALIGN", (0, 0), (-1, 0), "CENTER"), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
      ("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 8), ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#FDECEA"))]

def page(title, size, note=None):
    W, H = size
    def draw(c, doc):
        c.saveState()
        c.setStrokeColor(RED); c.setLineWidth(1.2)
        c.line(14*mm, H-18*mm, W/2-45*mm, H-18*mm); c.line(W/2+45*mm, H-18*mm, W-14*mm, H-18*mm)
        c.setFont("Helvetica-Bold", 20); c.setFillColor(colors.black); c.drawCentredString(W/2, H-20*mm, "Zidane Corporation")
        c.setFont("Helvetica-Bold", 15); c.setFillColor(RED); c.drawCentredString(W/2, H-28*mm, title)
        c.setFillColor(colors.black); y = H-35*mm
        for lbl, val in (("Reference:", REF), ("Currency:", "PKR"), ("Date:", DATE)):
            c.setFont("Helvetica-Bold", 9); c.drawString(14*mm, y, lbl)
            c.setFont("Helvetica", 9); c.drawString(14*mm + c.stringWidth(lbl, "Helvetica-Bold", 9) + 4, y, val); y -= 4.2*mm
        if note:
            c.setFont("Helvetica-Bold", 8); c.setFillColor(RED); c.drawRightString(W-14*mm, H-35*mm, note)
        c.setFillColor(colors.black); c.setFont("Helvetica", 8); c.drawCentredString(W/2, 9*mm, f"Page {doc.page}")
        c.restoreState()
    return draw

rs = lambda v: f"{v:,.0f}"
sell_total = sum(it["sell"] * it["qty"] for it in ITEMS)
buy_total = sum(it["high"] * it["qty"] for it in ITEMS)

data = [["SR#", "Item Description", "Brand", "UOM", "Qty", "Unit Rate (PKR)", "Amount (PKR)"]]
for it in ITEMS:
    data.append([str(it["sr"]), Paragraph(it["desc"], cell), Paragraph(it["shown"], cell), it["uom"], str(it["qty"]),
                 f"Rs {rs(it['sell'])}", f"Rs {rs(it['sell'] * it['qty'])}"])
data.append(["", "TOTAL", "", "", "", "", f"Rs {rs(sell_total)}"])
t = Table(data, repeatRows=1, colWidths=[10*mm, 72*mm, 34*mm, 16*mm, 10*mm, 22*mm, 24*mm])
t.setStyle(TableStyle(TS + [("ALIGN", (0, 1), (0, -1), "CENTER"), ("ALIGN", (3, 1), (4, -1), "CENTER"),
                            ("ALIGN", (5, 1), (6, -1), "RIGHT")]))
terms = ParagraphStyle("t", fontName="Helvetica", fontSize=8, leading=11)
story = [t, Spacer(1, 6*mm), Paragraph("<b>Terms:</b> Rates in PKR, taxes extra as applicable. Validity 15 days. "
         "Delivery: ex-stock items 3-5 working days; imported items 2-3 weeks, subject to availability.", terms)]
SimpleDocTemplate("Zidane_Quotation_Rajby_PR217.pdf", pagesize=A4, leftMargin=12*mm, rightMargin=12*mm, topMargin=50*mm,
                  bottomMargin=16*mm, title="Quotation Rajby PR# 217", author="Zidane Corporation"
                  ).build(story, onFirstPage=page("QUOTATION", A4), onLaterPages=page("QUOTATION", A4))

data = [["SR#", "Item Description", "Buy Brand", "Qty", "Buy Rate", "Buy Total", "Sell Rate", "Sell Total", "Profit", "Margin", "Src", "V/S/N"]]
for it in ITEMS:
    bt, st = it["high"] * it["qty"], it["sell"] * it["qty"]
    data.append([str(it["sr"]), Paragraph(it["desc"], cell), Paragraph(it["brand"], cell), str(it["qty"]), rs(it["high"]), rs(bt),
                 rs(it["sell"]), rs(st), rs(st - bt), f"{(st - bt) / st * 100:.1f}%", it["src"], it["status"]])
data.append(["", "GRAND TOTAL", "", "", "", rs(buy_total), "", rs(sell_total), rs(sell_total - buy_total),
             f"{(sell_total - buy_total) / sell_total * 100:.1f}%", "", ""])
t = Table(data, repeatRows=1, colWidths=[10*mm, 80*mm, 34*mm, 10*mm, 22*mm, 24*mm, 22*mm, 24*mm, 22*mm, 15*mm, 10*mm, 12*mm])
t.setStyle(TableStyle(TS + [("ALIGN", (0, 1), (0, -1), "CENTER"), ("ALIGN", (3, 1), (3, -1), "CENTER"),
                            ("ALIGN", (4, 1), (9, -1), "RIGHT"), ("ALIGN", (10, 1), (11, -1), "CENTER")]))
SimpleDocTemplate("Zidane_Buying_vs_Selling_Rajby_PR217.pdf", pagesize=landscape(A4), leftMargin=12*mm, rightMargin=12*mm,
                  topMargin=48*mm, bottomMargin=16*mm, title="Buying vs Selling Rajby PR# 217", author="Zidane Corporation"
                  ).build([t], onFirstPage=page("BUYING vs SELLING", landscape(A4), "INTERNAL - NOT FOR CUSTOMER"),
                          onLaterPages=page("BUYING vs SELLING", landscape(A4), "INTERNAL - NOT FOR CUSTOMER"))
print("buy", rs(buy_total), "sell", rs(sell_total), "profit", rs(sell_total - buy_total))
from collections import Counter
for k in "VSN":
    g = [it for it in ITEMS if it["status"] == k]
    print(k, len(g), rs(sum(i["sell"] * i["qty"] for i in g)))
