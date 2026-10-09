import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from items import ITEMS

wb = Workbook()
ws = wb.active; ws.title = "Rate Check"
F = "Arial"
hdr_fill = PatternFill("solid", fgColor="B71C1C")
thin = Side(style="thin", color="BFBFBF"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)

ws["A1"] = "Zidane Corporation - Quote PR-217 & PR-228: Brand + Buying Rate Check"; ws["A1"].font = Font(name=F, bold=True, size=14)
ws["A2"].font = Font(name=F, italic=True, size=9)
ws["A3"] = "Target markup on buying (High):"; ws["A3"].font = Font(name=F, bold=True)
ws["E3"] = 0.15; ws["E3"].number_format = "0%"; ws["E3"].font = Font(name=F, color="0000FF", bold=True)
ws["E3"].fill = PatternFill("solid", fgColor="FFFF00")
ws["F3"] = "<- yellow cell change karo, revised rates update ho jayenge"; ws["F3"].font = Font(name=F, size=9, italic=True)

heads = ["SR#", "Item Description", "Quoted Brand", "Qty", "Quoted Rate (PKR)", "Correct / Suggested Brand",
         "Buying Low (PKR)", "Buying High (PKR)", "Margin on High Buy", "Status", "Revised Rate (PKR)",
         "Quoted Total", "Revised Total", "Cost Total (High)", "Brand Changed?", "Confidence", "Source / Basis", "Note"]
HR = 5
for c, h in enumerate(heads, 1):
    cell = ws.cell(HR, c, h); cell.font = Font(name=F, bold=True, color="FFFFFF"); cell.fill = hdr_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = bd
ws.row_dimensions[HR].height = 32

r = HR + 1
first = r
for it in ITEMS:
    vals = [it["sr"], it["desc"], it["qbrand"], it["qty"], it["rate"], it["brand"], it["low"], it["high"]]
    for c, v in enumerate(vals, 1):
        ws.cell(r, c, v)
    ws.cell(r, 9, f"=IF(E{r}=0,0,(E{r}-H{r})/E{r})")
    ws.cell(r, 10, f'=IF(E{r}<G{r},"LOSS",IF(E{r}<H{r}*1.1,"THIN / RISK",IF(E{r}>H{r}*1.8,"OVERPRICED?","OK")))')
    ws.cell(r, 11, f"=IF(J{r}=\"OVERPRICED?\",ROUND(H{r}*1.35,-1),MAX(E{r},ROUND(H{r}*(1+$E$3),-1)))")
    ws.cell(r, 12, f"=D{r}*E{r}")
    ws.cell(r, 13, f"=D{r}*K{r}")
    ws.cell(r, 14, f"=D{r}*H{r}")
    ws.cell(r, 15, f'=IF(ISNUMBER(SEARCH(C{r},F{r})),"","YES")')
    ws.cell(r, 16, it["conf"]); ws.cell(r, 17, it["src"]); ws.cell(r, 18, it["note"])
    for c in range(1, 19):
        cell = ws.cell(r, c); cell.border = bd; cell.font = Font(name=F, size=9, color="0000FF" if c in (7, 8) else "000000")
        cell.alignment = Alignment(vertical="center", wrap_text=c in (2, 6, 17, 18))
    for c in (5, 7, 8, 11, 12, 13, 14):
        ws.cell(r, c).number_format = '#,##0;(#,##0);-'
    ws.cell(r, 9).number_format = "0%;-0%;-"
    r += 1
last = r - 1

ws.cell(r, 2, "TOTAL").font = Font(name=F, bold=True)
for c in (12, 13, 14):
    L = get_column_letter(c)
    cell = ws.cell(r, c, f"=SUM({L}{first}:{L}{last})"); cell.font = Font(name=F, bold=True); cell.number_format = "#,##0"
tot = r

red = PatternFill("solid", fgColor="F8CBAD"); amber = PatternFill("solid", fgColor="FFE699")
blue = PatternFill("solid", fgColor="BDD7EE"); green = PatternFill("solid", fgColor="C6EFCE")
rng = f"J{first}:J{last}"
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'J{first}="LOSS"'], fill=red))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'J{first}="THIN / RISK"'], fill=amber))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'J{first}="OVERPRICED?"'], fill=blue))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'J{first}="OK"'], fill=green))
ws.conditional_formatting.add(f"O{first}:O{last}", FormulaRule(formula=[f'O{first}="YES"'], fill=amber))

widths = [6, 42, 14, 6, 12, 26, 12, 12, 10, 13, 13, 14, 14, 14, 9, 8, 40, 36]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = ws.cell(HR + 1, 3)
ws.auto_filter.ref = f"A{HR}:R{last}"

# Summary
s = wb.create_sheet("Summary")
R = f"'Rate Check'!"
rows = [
    ("Total line items", f"=COUNTA({R}A{first}:A{last})"),
    ("LOSS items (quoted < lowest buying)", f'=COUNTIF({R}J{first}:J{last},"LOSS")'),
    ("THIN / RISK items (<10% over high buying)", f'=COUNTIF({R}J{first}:J{last},"THIN / RISK")'),
    ("OK items", f'=COUNTIF({R}J{first}:J{last},"OK")'),
    ("OVERPRICED? items (>80% over high buying)", f'=COUNTIF({R}J{first}:J{last},"OVERPRICED?")'),
    ("Items where brand must change", f'=COUNTIF({R}O{first}:O{last},"YES")'),
    ("", None),
    ("Quoted total (PKR)", f"={R}L{tot}"),
    ("Estimated cost at HIGH buying (PKR)", f"={R}N{tot}"),
    ("Gross margin on quote (PKR)", "=B10-B11"),
    ("Gross margin %", "=IF(B10=0,0,B12/B10)"),
    ("Loss on LOSS/THIN lines (PKR, at high buying)", f'=SUMPRODUCT(({R}E{first}:E{last}<{R}H{first}:H{last})*({R}H{first}:H{last}-{R}E{first}:E{last})*{R}D{first}:D{last})'),
    ("", None),
    ("Revised quote total (PKR)", f"={R}M{tot}"),
    ("Revised margin %", "=IF(B16=0,0,(B16-B11)/B16)"),
]
s["A1"] = "Summary - Quote PR-217 & PR-228"; s["A1"].font = Font(name=F, bold=True, size=13)
for i, (k, v) in enumerate(rows, 3):
    s.cell(i, 1, k).font = Font(name=F)
    if v: c = s.cell(i, 2, v); c.font = Font(name=F, bold=True); c.number_format = "#,##0"
for rr in (13, 17): s.cell(rr, 2).number_format = "0.0%"
s.column_dimensions["A"].width = 48; s.column_dimensions["B"].width = 18
# fix row refs: rows list starts at row 3, so B8->row10 etc. rebuild references

src = wb.create_sheet("Sources")
links = [
 ("toolsmart.pk - Licota collection", "https://www.toolsmart.pk/collections/licota"),
 ("toolsmart.pk - Licota 18in pipe wrench Rs 9,620", "https://www.toolsmart.pk/products/licota-18-heavy-duty-ridgid-type-pipe-wrench-cr-mo"),
 ("toolsmart.pk - Licota 9pc extra long hex", "https://www.toolsmart.pk/products/licota-9pcs-extra-long-type-hex-key-set"),
 ("toolsmart.pk - Licota 120pc socket set", "https://www.toolsmart.pk/products/licota-120pcs-1-4-3-8-1-2-dr-socket-set"),
 ("toolsmart.pk - grease guns (Licota 400cc Rs 7,760)", "https://www.toolsmart.pk/collections/all-tools/grease-gun"),
 ("kamadi.pk - Bosch GDS 18V-1000", "https://kamadi.pk/products/gds-18v-1000-bosch-cordless-impact-wrench-650n-m-18v"),
 ("kamadi.pk - Bosch GDX 180-Li", "https://kamadi.pk/products/gdx-180-li-bosch-cordless-impact-driver-wrench-180n-m-18v"),
 ("kamadi.pk - Bosch GWS 2200-230 H", "https://kamadi.pk/products/bosch-gws-2200-230-h-angle-grinder"),
 ("kamadi.pk - Insize 1205-300S", "https://kamadi.pk/products/insize-1205-300s-vernier-caliper-0-300mm"),
 ("kamadi.pk - Insize 1108-300", "https://kamadi.pk/products/1108-300-insize-digital-4-way-vernier-caliper-0-300mm"),
 ("kamadi.pk - Insize 3109-25A micrometer", "https://kamadi.pk/products/insize-3109-25a-digital-outside-micrometer-0-25mm"),
 ("ktools.pk - Bosch GDS 18V-1000 solo", "https://www.ktools.pk/products/bosch-gds-18v-1000-cordless-impact-wrench"),
 ("ktools.pk - Bosch GDX 180-Li", "https://www.ktools.pk/products/bosch-cordless-impact-wrench-gdx-180-li"),
 ("ktools.pk - Bosch GBH 180-Li", "https://www.ktools.pk/products/bosch-cordless-rotary-hammer-with-sds-plus-gbh-180-li"),
 ("ktools.pk - Bosch GSB 18V-150C", "https://www.ktools.pk/products/bosch-cordless-impact-drill-combi-gsb-18v-150-c"),
 ("ktools.pk - Bosch GWS 2200-230H", "https://www.ktools.pk/products/bosch-gws-2200-230h-angle-grinder-230mm-2200w"),
 ("amjadhardware.com - Insize 1108-300", "https://amjadhardware.com/product/1108-300-digital-vernier-caliper-0-300mm-0-12-insize/"),
 ("purchaser.com.pk - Insize 1205-200S", "https://purchaser.com.pk/insize-4-way-vernier-caliper-1205-200s-en/"),
 ("powerhouseexpress.com.pk - grinders", "https://powerhouseexpress.com.pk/collections/grinders"),
 ("rafiqbrothers.com - hand pallet trucks", "https://www.rafiqbrothers.com/en/category/hand-pallet-trucks/"),
 ("eqmachines.com - hand pallet truck", "https://eqmachines.com/products/hand-pallet-truck-pakistan"),
 ("pakwheels.com - bottle jacks", "https://www.pakwheels.com/accessories-spare-parts/car-jack/213239"),
 ("ingcotool.pk - Total 8in bench vise", "https://ingcotool.pk/product/total-swivel-base-bench-vise-with-anvil-6-150mm-tht6166/"),
 ("naheed.pk - WD-40 330ml Rs 1,380", "https://www.naheed.pk/wd-40-330ml"),
 ("autohub.pk - WD-40 330ml Rs 1,299", "https://autohub.pk/products/wd-40-330ml"),
 ("imartpk.com - Bosch GWS 22-230 H", "https://imartpk.com/products/gws-22-230-h-professional"),
 ("goldentools.ae - Licota 3-jaw puller (UAE ref)", "https://www.goldentools.ae/gttecom/singleitemmob/LICOTA/Bearing-and-Seal-Tools/?brcd=4712834507742"),
 ("amazon.ae - Licota 6in 3-jaw puller (UAE ref)", "https://www.amazon.ae/LICOTA-GEAR-PULLER-ATB-1002C-AT-2102A/dp/B08CTD3842"),
 ("goldentools.ae - Licota 12mm hex key (UAE ref)", "https://www.goldentools.ae/item/4712818539608"),
 ("goldpeaktools.com.ph - Licota AWX-2603GN 7-drawer (ref)", "https://shop.goldpeaktools.com.ph/products/licota-awx-2603gn-tool-cabinet-carriage-7-drawers"),
]
src["A1"] = "Sources (checked Oct 2026 - prices change, confirm before PO)"; src["A1"].font = Font(name=F, bold=True)
for i,(a,b) in enumerate(links,3):
    src.cell(i,1,a).font=Font(name=F); c=src.cell(i,2,b); c.hyperlink=b; c.font=Font(name=F,color="0563C1",underline="single")
src.column_dimensions["A"].width=55; src.column_dimensions["B"].width=95
wb.save(sys.argv[1])
