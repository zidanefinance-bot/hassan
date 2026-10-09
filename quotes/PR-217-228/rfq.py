"""RFQ sheets for items that are not (exactly) listed online - send to dealers, fill rates."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from apply_listings import ITEMS_ONLINE

F = "Arial"
thin = Side(style="thin", color="BFBFBF"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
TOOLS = ("Yato", "Dormer", "YG-1", "Vertex", "Insize", "Masada", "Taiwan", "Bosch", "Generic (Taiwan)")

def group(it):
    b = it["brand"]
    if b.startswith("Licota"): return "Licota Dealer"
    if b.startswith(TOOLS): return "Tool Importer"
    return "Welding-Safety-Consumables"

wb = Workbook(); wb.remove(wb.active)
items = [i for i in ITEMS_ONLINE if i["conf"] in ("N", "S")]
for g in ("Licota Dealer", "Tool Importer", "Welding-Safety-Consumables"):
    ws = wb.create_sheet(g)
    ws["A1"] = f"Zidane Corporation - Request for Quotation ({g})"; ws["A1"].font = Font(name=F, bold=True, size=13)
    ws["A2"] = "Ref: PR-217 & PR-228 | Please fill Unit Rate (PKR), brand offered and delivery time. Yellow cells = aap ne bharne hain."
    ws["A2"].font = Font(name=F, size=9, italic=True)
    hd = ["SR#", "Item Description", "Qty", "Preferred Brand", "Brand Offered", "Unit Rate (PKR)", "Delivery (days)", "Remarks"]
    for c, h in enumerate(hd, 1):
        x = ws.cell(4, c, h); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = PatternFill("solid", fgColor="B71C1C"); x.border = bd
        x.alignment = Alignment(horizontal="center")
    r = 5
    for it in [i for i in items if group(i) == g]:
        pref = it["brand"].split(" / ")[0]
        for c, v in enumerate((it["sr"], it["desc"], it["qty"], pref, None, None, None, None), 1):
            x = ws.cell(r, c, v); x.font = Font(name=F, size=10); x.border = bd
        for c in (5, 6, 7):
            ws.cell(r, c).fill = PatternFill("solid", fgColor="FFF2CC")
        ws.cell(r, 6).number_format = "#,##0"
        r += 1
    for col, w in zip("ABCDEFGH", (7, 48, 7, 22, 18, 16, 14, 24)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    print(g, r - 5)
wb.save("RFQ_PR217_PR228.xlsx")
