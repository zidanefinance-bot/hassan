from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Spacer
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from decimal import Decimal, ROUND_HALF_UP
from apply_listings import ITEMS_ONLINE as ITEMS
from overrides import SELL_OVERRIDE

MARKUP = 0.15
RED = colors.HexColor("#B71C1C")
KNOWN = ["Bosch Professional", "DeWalt", "Draper", "Vertex (Taiwan)", "Toho (Taiwan)", "Masada (Japan)", "Licota", "Yato", "Insize", "Dormer",
         "YG-1", "Norton", "Karam", "WD-40", "Industrial", "Bosch"]
MAP = {"Generic (Taiwan)": "Taiwan", "Taiwan HSS": "Taiwan", "Harris type": "Industrial", "Local": "Local", "Generic": "Taiwan"}

def brand(b):
    if b in DIRECT: return DIRECT[b]
    if b.startswith(DORMER): return DORMER
    first = b.split(" / ")[0].strip()
    for k, v in MAP.items():
        if first.startswith(k): return v
    for k in KNOWN:
        if first.startswith(k): return k
    raise ValueError(b)

DIRECT = {"Harden": "Harden", "Total": "Total", "Ingco": "Ingco", "Tolsen": "Tolsen", "Ronix": "Ronix", "Tanaka": "Tanaka",
          "Blue Eagle": "Blue Eagle", "Leather Forge": "Leather Forge", "7CF": "7CF", "Mubah": "Mubah", "Zodian": "Zodian",
          "Fortrex": "Fortrex", "WM": "WM", "Industrial (VMAX)": "VMAX", "WACTA20": "Industrial", "Rhodius": "Rhodius", "Generic HSS": "HSS", "SMT": "SMT", "Stainarc": "Stainarc", "Taiwan": "Taiwan", "JET (Taiwan)": "JET (Taiwan)"}
DORMER = "Dormer"

def r10(x):
    return float(Decimal(str(round(x, 6))).quantize(Decimal("1E1"), rounding=ROUND_HALF_UP))

def rate(it):
    if it["sr"] in SELL_OVERRIDE: return SELL_OVERRIDE[it["sr"]]
    hi, q = it["high"], it["rate"]
    if q > hi * 1.8: return r10(hi * 1.35)
    return max(q, r10(hi * (1 + MARKUP)))

KEEP = ("Bosch", "WD-40", "Insize")
FOREIGN = {int(l) for l in open("foreign_sourced.txt") if l.strip() and not l.startswith("#")}
def shown(it):
    b = brand(it["brand"])
    if it["qbrand"] == "Industrial" and not b.startswith(KEEP):
        b = "Industrial"
    return b + " / Equivalent" if it["sr"] in FOREIGN else b

rows = [(it["sr"], it["desc"], shown(it), it["qty"], rate(it)) for it in ITEMS]

# ---- PDF in the same layout as the original quotation ----
PER_PAGE = 46
pages = [rows[i:i + PER_PAGE] for i in range(0, len(rows), PER_PAGE)]
W, H = A4

def header_footer(c, doc):
    c.saveState()
    c.setStrokeColor(RED); c.setLineWidth(1.2)
    c.line(14*mm, H-18*mm, 68*mm, H-18*mm); c.line(W-68*mm, H-18*mm, W-14*mm, H-18*mm)
    c.setFont("Helvetica-Bold", 20); c.setFillColor(colors.black); c.drawCentredString(W/2, H-20*mm, "Zidane Corporation")
    c.setFont("Helvetica-Bold", 16); c.setFillColor(RED); c.drawCentredString(W/2, H-29*mm, "REVISED QUOTATION")
    c.setFillColor(colors.black); c.setFont("Helvetica-Bold", 9)
    y = H-35*mm
    for lbl, val in (("Reference:", "PR-217 & PR-228 (Rev-1)"), ("Currency:", "PKR"), ("Date:", "09-Oct-2026")):
        c.setFont("Helvetica-Bold", 9); c.drawString(14*mm, y, lbl)
        c.setFont("Helvetica", 9); c.drawString(14*mm + c.stringWidth(lbl, "Helvetica-Bold", 9) + 4, y, val); y -= 4.2*mm
    c.setFont("Helvetica", 9); c.drawCentredString(W/2, 10*mm, f"Page {doc.page} of {len(pages)}")
    c.restoreState()

doc = SimpleDocTemplate("Zidane_Revised_Quotation_PR217_PR228.pdf", pagesize=A4, leftMargin=12*mm, rightMargin=12*mm,
                        topMargin=50*mm, bottomMargin=16*mm, title="Revised Quotation PR-217 & PR-228", author="Zidane Corporation")
story = []
for i, pg in enumerate(pages):
    data = [["SR#", "Item Description", "Brand", "Qty", "Unit Rate (PKR)"]]
    data += [[str(s), d, b, str(q), f"Rs {r:,.0f}"] for s, d, b, q, r in pg]
    t = Table(data, colWidths=[13*mm, 85*mm, 35*mm, 16*mm, 37*mm], repeatRows=1, rowHeights=[7*mm] + [4.7*mm]*len(pg))
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), RED), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9), ("FONT", (0, 1), (-1, -1), "Helvetica", 7.4),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"), ("ALIGN", (0, 1), (0, -1), "CENTER"), ("ALIGN", (2, 1), (3, -1), "CENTER"),
        ("ALIGN", (4, 1), (4, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5), ("TOPPADDING", (0, 0), (-1, -1), 1.5),
    ]))
    story.append(t)
    if i < len(pages) - 1:
        from reportlab.platypus import PageBreak; story.append(PageBreak())
doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)

# ---- Editable Excel copy ----
wb = Workbook(); ws = wb.active; ws.title = "Revised Quotation"
F = "Arial"; thin = Side(style="thin", color="BFBFBF"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
ws["A1"] = "Zidane Corporation - REVISED QUOTATION"; ws["A1"].font = Font(name=F, bold=True, size=14)
ws["A2"] = "Reference: PR-217 & PR-228 (Rev-1)    Currency: PKR    Date: 09-Oct-2026"; ws["A2"].font = Font(name=F)
hd = ["SR#", "Item Description", "Brand", "Qty", "Unit Rate (PKR)", "Amount (PKR)"]
for c, h in enumerate(hd, 1):
    x = ws.cell(4, c, h); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = PatternFill("solid", fgColor="B71C1C")
    x.alignment = Alignment(horizontal="center"); x.border = bd
r = 5
for s, d, b, q, rt in rows:
    for c, v in enumerate((s, d, b, q, rt), 1): ws.cell(r, c, v)
    ws.cell(r, 6, f"=D{r}*E{r}")
    for c in range(1, 7):
        x = ws.cell(r, c); x.font = Font(name=F, size=10); x.border = bd
    ws.cell(r, 5).number_format = ws.cell(r, 6).number_format = "#,##0"
    r += 1
ws.cell(r, 5, "TOTAL").font = Font(name=F, bold=True)
x = ws.cell(r, 6, f"=SUM(F5:F{r-1})"); x.font = Font(name=F, bold=True); x.number_format = "#,##0"
for col, w in zip("ABCDEF", (7, 50, 20, 7, 16, 18)): ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
wb.save("Zidane_Revised_Quotation_PR217_PR228.xlsx")
print(len(rows), len(pages), sum(q*rt for _,_,_,q,rt in rows))
from collections import Counter; print(Counter(b for _,_,b,_,_ in rows))
