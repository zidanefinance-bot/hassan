"""Sheet for Hassan: unverified items (status N/S) to fill with rate + source link."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from apply_listings import ITEMS_ONLINE

F = "Arial"
thin = Side(style="thin", color="BFBFBF"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
items = sorted([i for i in ITEMS_ONLINE if i["conf"] in ("N", "S")], key=lambda i: -i["qty"] * i["high"])
wb = Workbook(); ws = wb.active; ws.title = "Rates Needed"
ws["A1"] = "Rates Needed - PR-217 & PR-228 (unverified items, highest value first)"; ws["A1"].font = Font(name=F, bold=True, size=13)
ws["A2"] = "Peeli cells bharo: Unit Rate (PKR) aur Source Link. Example row 3 dekho."; ws["A2"].font = Font(name=F, size=9, italic=True)
hd = ["#", "SR#", "Item Description", "Qty", "Brand Required", "Value at Stake (PKR)", "Unit Rate (PKR)", "Source Link", "Status"]
ex = ["eg", 8, "7 DRAWER TOOL TROLLEY", 3, "Licota", None, 272440, "https://www.toolsmart.pk/... (Licota 7 Drawer Cabinet)", "example"]
for c, h in enumerate(hd, 1):
    x = ws.cell(4, c, h); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = PatternFill("solid", fgColor="B71C1C"); x.border = bd
for c, v in enumerate(ex, 1):
    x = ws.cell(3, c, v); x.font = Font(name=F, size=9, italic=True, color="808080")
r = 5
for n, it in enumerate(items, 1):
    vals = (n, it["sr"], it["desc"], it["qty"], it["brand"].split(" / ")[0], round(it["qty"] * it["high"]), None, None,
            "Spec farq" if it["conf"] == "S" else "Online nahi mila")
    for c, v in enumerate(vals, 1):
        x = ws.cell(r, c, v); x.font = Font(name=F, size=10); x.border = bd
    ws.cell(r, 6).number_format = "#,##0"; ws.cell(r, 7).number_format = "#,##0"
    for c in (7, 8):
        ws.cell(r, c).fill = PatternFill("solid", fgColor="FFF2CC")
    r += 1
ws.cell(r, 5, "TOTAL").font = Font(name=F, bold=True)
x = ws.cell(r, 6, f"=SUM(F5:F{r-1})"); x.font = Font(name=F, bold=True); x.number_format = "#,##0"
for col, w in zip("ABCDEFGHI", (5, 7, 44, 6, 26, 16, 16, 45, 16)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"
wb.save("Rates_Needed_PR217_PR228.xlsx")
print(len(items))
from collections import Counter
print(Counter(i["brand"].split()[0] for i in items))
