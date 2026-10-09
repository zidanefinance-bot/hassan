"""Internal buying vs selling sheet (PDF) for PR-217 & PR-228 Rev-1.
Reads recalculated values from the Rate Check workbook; run after build.py + revised_quote.py + recalc."""
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Spacer, Paragraph
from reportlab.lib.styles import ParagraphStyle
from openpyxl import load_workbook

RED = colors.HexColor("#B71C1C")
FOREIGN = {int(l) for l in open("foreign_sourced.txt") if l.strip() and not l.startswith("#")}

ws = load_workbook("Quote_PR217_PR228_Rate_Check.xlsx", data_only=True)["Rate Check"]
rows = []
for r in ws.iter_rows(min_row=6, values_only=True):
    if not isinstance(r[0], int):
        continue
    sr, desc, qty, brand, buy, sell, status = r[0], r[1], r[3], r[5], r[7], r[10], r[15]
    rows.append(dict(sr=sr, desc=desc, brand=brand, qty=qty, buy=buy, sell=sell,
                     bt=buy * qty, st=sell * qty, src="Intl" if sr in FOREIGN else "PK", vs=status))

BT = sum(x["bt"] for x in rows); ST = sum(x["st"] for x in rows); PR = ST - BT

def rs(v): return f"{v:,.0f}"
def pct(p, s): return f"{p / s * 100:.1f}%" if s else "-"

def grp(key):
    out = []
    for k in sorted({x[key] for x in rows}):
        g = [x for x in rows if x[key] == k]
        b = sum(x["bt"] for x in g); s = sum(x["st"] for x in g)
        out.append([k, str(len(g)), rs(b), rs(s), rs(s - b), pct(s - b, s)])
    return out

W, H = landscape(A4)

def header_footer(c, doc):
    c.saveState()
    c.setFont("Helvetica-Bold", 15); c.drawString(10*mm, H-12*mm, "Zidane Corporation")
    c.setFillColor(RED); c.setFont("Helvetica-Bold", 12)
    c.drawString(10*mm, H-18*mm, "BUYING vs SELLING  -  PR-217 & PR-228 (Rev-1)")
    c.setFillColor(colors.black); c.setFont("Helvetica", 8)
    c.drawRightString(W-10*mm, H-12*mm, "INTERNAL - NOT FOR CUSTOMER")
    c.drawRightString(W-10*mm, H-18*mm, "Date: 09-Oct-2026   |   Currency: PKR")
    c.setStrokeColor(RED); c.setLineWidth(1); c.line(10*mm, H-21*mm, W-10*mm, H-21*mm)
    c.setFont("Helvetica", 7.5)
    c.drawString(10*mm, 7*mm, "Buy = highest verified online rate. Src: PK = Pakistani seller, Intl = UAE/UK/US/EU rate converted (AED x90 shelf basis). "
                 "V = exact listing, S = size/pack/brand differs.")
    c.drawRightString(W-10*mm, 7*mm, f"Page {doc.page}")
    c.restoreState()

TS = [("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.6), ("BACKGROUND", (0, 0), (-1, 0), RED),
      ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONT", (0, 1), (-1, -1), "Helvetica", 7.4),
      ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
      ("TOPPADDING", (0, 0), (-1, -1), 1.3), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.3)]

story = []
h = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=10, spaceAfter=3)
story.append(Paragraph("Summary", h))
summ = [["", "Items", "Buying (PKR)", "Selling (PKR)", "Profit (PKR)", "Margin"],
        ["TOTAL", str(len(rows)), rs(BT), rs(ST), rs(PR), pct(PR, ST)]]
summ += [["Status " + r[0]] + r[1:] for r in grp("vs")]
summ += [["Source " + r[0]] + r[1:] for r in grp("src")]
t = Table(summ, colWidths=[30*mm, 16*mm, 32*mm, 32*mm, 32*mm, 18*mm], hAlign="LEFT")
t.setStyle(TableStyle(TS + [("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("FONT", (0, 1), (-1, 1), "Helvetica-Bold", 7.6),
                             ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#FDECEA"))]))
story += [t, Spacer(1, 5*mm), Paragraph("Line items", h)]

data = [["SR#", "Item Description", "Buy Brand", "Qty", "Buy Rate", "Buy Total", "Sell Rate", "Sell Total",
         "Profit", "Margin", "Src", "V/S"]]
for x in rows:
    p = x["st"] - x["bt"]
    data.append([str(x["sr"]), x["desc"][:48], str(x["brand"])[:24], str(x["qty"]), rs(x["buy"]), rs(x["bt"]),
                 rs(x["sell"]), rs(x["st"]), rs(p), pct(p, x["st"]), x["src"], x["vs"]])
data.append(["", "GRAND TOTAL", "", "", "", rs(BT), "", rs(ST), rs(PR), pct(PR, ST), "", ""])
t = Table(data, repeatRows=1, colWidths=[11*mm, 72*mm, 38*mm, 11*mm, 21*mm, 24*mm, 21*mm, 24*mm, 22*mm, 14*mm, 10*mm, 9*mm])
t.setStyle(TableStyle(TS + [
    ("ALIGN", (0, 0), (0, -1), "CENTER"), ("ALIGN", (3, 0), (3, -1), "CENTER"), ("ALIGN", (4, 0), (9, -1), "RIGHT"),
    ("ALIGN", (10, 0), (11, -1), "CENTER"), ("ALIGN", (0, 0), (-1, 0), "CENTER"),
    ("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 7.6), ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#FDECEA")),
    ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F7F7F7")])]))
story.append(t)

doc = SimpleDocTemplate("Zidane_Buying_vs_Selling_PR217_PR228.pdf", pagesize=landscape(A4), leftMargin=10*mm,
                        rightMargin=10*mm, topMargin=25*mm, bottomMargin=12*mm,
                        title="Buying vs Selling PR-217 & PR-228", author="Zidane Corporation")
doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
print(len(rows), rs(BT), rs(ST), rs(PR), pct(PR, ST))
