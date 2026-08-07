#!/usr/bin/env python3
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HDR  = "1F4E79"
MED  = "2E75B6"
LGRY = "F2F2F2"
thin = Side(style='thin', color="AAAAAA")
thick = Side(style='medium', color="000000")
bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)
hbdr = Border(left=thick, right=thick, top=thick, bottom=thick)

def style(c, bg=None, fc="000000", bold=False, size=10, halign="center", valign="center", wrap=False, border=bdr):
    if bg:
        c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    c.font = Font(color=fc, bold=bold, size=size)
    c.alignment = Alignment(horizontal=halign, vertical=valign, wrap_text=wrap)
    c.border = border

def make_invoice(wb, po_num, date, unit, items):
    """
    items = list of (item_name, uom, qty, rate, tax_pct)
    tax_pct used only to show amount — invoice shows excl tax + tax + total
    """
    ws = wb.create_sheet(title=f"PO#{po_num}")
    ws.sheet_view.showGridLines = False

    # Column widths: A=5, B=45, C=8, D=8, E=14, F=14, G=10, H=16
    col_w = [5, 45, 8, 8, 14, 14, 10, 16]
    for i, w in enumerate(col_w, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── COMPANY HEADER ──
    ws.merge_cells("A1:H1")
    c = ws["A1"]; c.value = "ZIDANE CORPORATION"
    style(c, bg=HDR, fc="FFFFFF", bold=True, size=16, border=hbdr)
    ws.row_dimensions[1].height = 38

    ws.merge_cells("A2:H2")
    c = ws["A2"]; c.value = "B-55, PNT Cooperative Society, Near Shaan Circle, Karachi  |  NTN: 5795308-1"
    style(c, bg=MED, fc="FFFFFF", size=9, border=hbdr)
    ws.row_dimensions[2].height = 18

    ws.merge_cells("A3:H3")
    c = ws["A3"]; c.value = "Email: zidanecorporation@gmail.com"
    style(c, bg=MED, fc="FFFFFF", size=9, border=hbdr)
    ws.row_dimensions[3].height = 15

    # ── INVOICE TITLE ──
    ws.row_dimensions[4].height = 6

    ws.merge_cells("A5:H5")
    c = ws["A5"]; c.value = "INVOICE"
    style(c, bg="FFFFFF", fc=HDR, bold=True, size=18, border=Border())
    ws.row_dimensions[5].height = 32

    # ── BILL TO / PO INFO ──
    ws.row_dimensions[6].height = 5

    # Left block — Bill To
    ws.merge_cells("A7:D7")
    c = ws["A7"]; c.value = "BILL TO:"
    style(c, bg=LGRY, fc=HDR, bold=True, size=9, halign="left", border=Border())

    ws.merge_cells("A8:D8")
    c = ws["A8"]; c.value = "Rajby Industries (Private) Limited"
    style(c, fc="000000", bold=True, size=11, halign="left", border=Border())

    ws.merge_cells("A9:D9")
    c = ws["A9"]; c.value = "Plot #38/39 & 77/78, Sector 27, Korangi Industrial Area, Karachi"
    style(c, fc="444444", size=9, halign="left", wrap=True, border=Border())
    ws.row_dimensions[9].height = 24

    for r in [7, 8]:
        ws.row_dimensions[r].height = 18

    # Right block — Invoice details
    detail_labels = ["PO #:", "Invoice Date:", "Operating Unit:", "Payment Terms:"]
    detail_values = [po_num, date, unit, "As per PO"]

    for i, (lbl, val) in enumerate(zip(detail_labels, detail_values)):
        r = 7 + i
        ws.merge_cells(f"E{r}:F{r}")
        c = ws.cell(row=r, column=5, value=lbl)
        style(c, bg=LGRY, fc="444444", bold=True, size=9, halign="right", border=Border())
        ws.merge_cells(f"G{r}:H{r}")
        c = ws.cell(row=r, column=7, value=val)
        bold_val = (lbl == "PO #:")
        style(c, bg="FFFFFF", fc=HDR if bold_val else "000000", bold=bold_val, size=10 if bold_val else 9,
              halign="left", border=Border())
        ws.row_dimensions[r].height = 18

    ws.row_dimensions[10].height = 12

    # ── ITEMS TABLE HEADER ──
    row = 11
    headers = ["Sr#", "Item Description", "UOM", "Qty", "Rate (PKR)", "Amount (PKR)", "Tax %", "Total incl. Tax"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        style(c, bg=HDR, fc="FFFFFF", bold=True, size=10, border=hbdr)
    ws.row_dimensions[row].height = 28
    ws.freeze_panes = f"A{row+1}"
    row += 1

    # ── ITEMS ──
    grand_excl = grand_tax = grand_incl = 0
    for sr, (name, uom, qty, rate, tax_pct) in enumerate(items, 1):
        amt   = qty * rate
        tax   = round(amt * tax_pct / 100, 2)
        total = amt + tax
        grand_excl += amt; grand_tax += tax; grand_incl += total

        bg = "FFFFFF" if sr % 2 == 1 else LGRY
        vals = [sr, name, uom, qty, rate, amt, f"{tax_pct}%" if tax_pct else "—", total]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=v)
            halign = "left" if col == 2 else "center"
            style(c, bg=bg, halign=halign, wrap=(col==2), border=bdr)
            if col in [5, 6, 8] and isinstance(v, (int, float)):
                c.number_format = '#,##0.00'
        ws.row_dimensions[row].height = 20
        row += 1

    # ── TOTALS ──
    row += 1
    def total_row(lbl, val, bg_lbl=LGRY, bg_val="FFFFFF", bold=False):
        nonlocal row
        ws.merge_cells(f"A{row}:F{row}")
        c = ws.cell(row=row, column=1, value=lbl)
        style(c, bg=bg_lbl, halign="right", bold=bold, size=10, border=bdr)
        ws.merge_cells(f"G{row}:H{row}")
        c = ws.cell(row=row, column=7, value=val)
        sz = 11 if bold else 10
        style(c, bg=bg_val, fc="000000" if not bold else "FFFFFF", bold=bold, size=sz, halign="center", border=hbdr if bold else bdr)
        c.number_format = '#,##0.00'
        ws.row_dimensions[row].height = 20
        row += 1

    total_row("Sub Total (Excl. Tax):", grand_excl)
    total_row("Tax Amount:", grand_tax)
    total_row("GRAND TOTAL:", grand_incl, bg_lbl=HDR, bg_val=HDR, bold=True)

    # Update last row font colors for grand total
    ws.cell(row=row-1, column=1).font = Font(color="FFFFFF", bold=True, size=10)
    ws.cell(row=row-1, column=7).font = Font(color="FFFFFF", bold=True, size=11)

    row += 1
    ws.row_dimensions[row].height = 8
    row += 1

    # ── SIGNATURE BLOCK ──
    ws.merge_cells(f"A{row}:H{row}")
    c = ws.cell(row=row, column=1, value="Prepared By: Hassan  |  Authorized Signatory: ___________________  |  Date: ___________")
    style(c, fc="666666", size=9, halign="left", border=Border())
    ws.row_dimensions[row].height = 22

    return ws

# ── PO DATA ──
# (item_name, uom, qty, rate, tax_pct)

pos = {
    "1183": {
        "date": "02-Jul-2026", "unit": "Washing PVT",
        "items": [
            ("Pyodine 60ml",           "PCS", 30, 225,  0),
            ("Cotton Medicated 25grm", "PKT", 20, 20,   0),
            ('Bandages 3"',            "PCS", 60, 25,   0),
            ('Bandages 4"',            "PCS", 60, 30,   0),
            ('Bandages 5"',            "PCS", 60, 34,   0),
            ("Tourniquet",             "PCS", 20, 200,  0),
            ("Ointment (Polyfax)",     "PCS", 20, 240,  0),
            ("Betnesol 7.5ml",         "PCS", 20, 230,  0),
            ("Dettol Bottle 50ml",     "PCS", 20, 260,  0),
            ("Paratule",               "PKT", 15, 320,  0),
        ]
    },
    "17427": {
        "date": "02-Jul-2026", "unit": "Cut to Pack PVT",
        "items": [
            ("Pyodine 60ml",                 "PCS", 40, 233,  5),
            ('Bandages 2"',                  "PCS", 36, 25,   5),
            ('Bandages 3"',                  "PCS", 36, 34,   5),
            ("Saniplast",                    "PKT", 12, 357,  5),
            ("Crepe Bandage",                "PCS", 24, 49,   5),
            ("Surgical Cotton",              "PCS", 24, 118,  5),
            ("Poly Flex",                    "PCS", 12, 250,  5),
            ("Burnol 20gm (Tube)",           "PCS", 12, 123,  5),
            ("Iodex Mini Jar",               "PCS", 20, 123,  5),
            ("Surgical Tape",                "PCS", 12, 93,   5),
            ("Benzoin Compound Tincture",    "PCS", 24, 82,   5),
            ("Dettol Bottle 50ml",           "PCS", 24, 270,  5),
            ("Surgical Safety Pin",          "PKT", 4,  49,   5),
        ]
    },
    "17412": {
        "date": "02-Jul-2026", "unit": "Cut to Pack PVT",
        "items": [
            ("Rice Spoon",               "PCS", 24, 50,  0),
            ('Ceramic Plates White 8"',  "PCS", 6,  295, 5),
            ('Ceramic Plates White 12"', "PCS", 6,  300, 5),
            ("Glass Jug 1.5 Ltr",        "PCS", 2,  700, 5),
            ("Deli Glass 6 pcs (Set)",   "SET", 2,  800, 5),
        ]
    },
    "17426": {
        "date": "02-Jul-2026", "unit": "Cut to Pack PVT",
        "items": [
            ("Plastic Jar (Small Size)", "PCS", 36, 70, 5),
        ]
    },
    "1163": {
        "date": "30-Jun-2026", "unit": "Washing PVT",
        "items": [
            ("Kafgeer 1 KG (Steel)", "PCS", 2, 1000, 0),
            ("Kafgeer 2 KG (Steel)", "PCS", 3, 3000, 0),
        ]
    },
    "17257": {
        "date": "29-Jun-2026", "unit": "Cut to Pack PVT",
        "items": [
            ("Air Freshner Room Spray 300ml Topic Brand","PCS", 96,  260,   5),
            ("Bleach 30KG Local",                        "CAN", 14,  970,   5),
            ("Broom Hard",                               "PCS", 200, 235,   5),
            ("Broom Soft PP",                            "PCS", 200, 200,   0),
            ("China Brush",                              "PCS", 60,  300,   5),
            ("Dustbin Bag 18x24 Black",                  "KG",  50,  300,   5),
            ("Dustbin Bag 30x50 Black",                  "KG",  50,  300,   5),
            ("Duster Fabric",                            "PCS", 96,  37,    0),
            ("Glass Cleaner Glint 500ml",                "BTL", 96,  250,   0),
            ("Hand Brush",                               "PCS", 12,  300,   0),
            ("Hand Wash Lose Local",                     "KG",  200, 270,   5),
            ("Handwash Lifeboy 140ml",                   "BTL", 48,  300,   0),
            ("Harpic 500ml",                             "BTL", 96,  380,   0),
            ("Lux Soap 70 Grams",                        "PCS", 100, 150,   5),
            ("Mop Big 18x18 Wooden",                     "PCS", 100, 394,   0),
            ("Mop Refill (600 Gram)",                    "PCS", 200, 347,   5),
            ("Mop Stick",                                "PCS", 60,  473,   5),
            ("Mortein Spray",                            "PCS", 48,  590,   5),
            ("Phenyl Zip Local 3 Ltr",                   "BTL", 200, 420,   0),
            ("Roomi Box",                                "CTN", 2,   18000, 5),
            ("Safety Match 10 PCS",                      "PKT", 15,  100,   5),
            ("Scotch Bright",                            "PCS", 24,  70,    5),
            ("Scraper Plastic Supri",                    "PCS", 60,  105,   5),
            ("Tile Wash Sweep Local 30KG",               "CAN", 7,   1180,  0),
            ("Tissue Box Rose Petal Pop Up 150 Sheet",   "PKT", 150, 270,   5),
            ("Tissue Paper Roll",                        "PCS", 300, 120,   5),
            ("VIM Toilet Bleach Powder Lose Local",      "PCS", 25,  180,   0),
            ("Viper Large 4ft",                          "PCS", 50,  579,   0),
        ]
    },
    "17308": {
        "date": "30-Jun-2026", "unit": "Cut to Pack PVT",
        "items": [
            ("Coffee (100g)",                    "PCS", 3,   1600, 0),
            ("Every Day Tea Whitener 2KG Pouch", "PKT", 260, 4185, 0),
            ("Lipton Lemon Green Tea 100 Bags",  "PKT", 30,  1150, 0),
            ("Lipton Tea Bag 600 Bags",          "BOX", 11,  3300, 0),
            ("Nescafe (200gm)",                  "PCS", 18,  2800, 0),
            ("Nestle Water 500ml",               "CRT", 23,  680,  0),
            ("Sugar 2 KGS Pack",                 "PKT", 350, 305,  0),
            ("Tapal Family Mixture (900gm)",      "PKT", 176, 1850, 0),
        ]
    },
    "1194": {
        "date": "06-Jul-2026", "unit": "Washing PVT",
        "items": [
            ("Molty Foam Sheet (3'x6'x2\") 980", "SHT", 6, 3579, 0),
        ]
    },
    "406": {
        "date": "03-Jul-2026", "unit": "Utilities PVT",
        "items": [
            ("Ceramic Fibre Rope 10mm Dia x 30m (Temp 150-200°C)", "ROL", 4, 15789, 0),
        ]
    },
}

wb = Workbook()
wb.remove(wb.active)  # remove default sheet

for po_num, data in pos.items():
    make_invoice(wb, po_num, data["date"], data["unit"], data["items"])

wb.save("Zidane_Invoices_Rajby.xlsx")
print("Done! 9 invoices created.")
